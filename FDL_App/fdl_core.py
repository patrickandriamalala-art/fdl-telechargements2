# -*- coding: utf-8 -*-
"""
FDL - logique de l'application (aucune interface ici, donc testable).

Projet = dictionnaire enregistré en JSON :
  debut, fin           : dates ISO (AAAA-MM-JJ) de la période à planifier
  ordre                : ordre de rotation des groupes
  premier_groupe       : groupe du 1er passage (Fidirana 1° du 1er dimanche)
  tokana               : dernier dimanche du mois = Fidirana tokana (1 seul groupe)
  role5                : utiliser aussi le rôle 5 (Vavaka fangatahana)
  groupes              : {groupe: {"membres": [{"nom": str, "absences": [[debut, fin], ...]}]}}
  programme            : [{"date": iso, "groupes": [g1, g2]}]
  affectations         : {"iso|groupe": {"1": nom, "2": nom, "3": nom, "4": nom, "5": nom}}
  verrous              : clés d'affectations modifiées à la main (conservées à la régénération)
  historique           : {groupe: {nom: {"n", "roles", "dernier_role", "depuis"}}}  (périodes précédentes)

Deux membres de groupes différents peuvent porter le même nom : ce sont des
personnes différentes (l'identité d'un membre = groupe + nom).
"""
import copy
import datetime as dt
import itertools
import json
import os
import sys

GROUPES_DEFAUT = ["ISAIA", "ESTERA", "JAONA", "ROTA", "KEFASY"]
ROLES_LIBELLES = {
    1: "Vavaka vakiteny (VVK VKT)",
    2: "Vakiteny I (VKT 1)",
    3: "Vakiteny II (VKT 2)",
    4: "Vakiteny III (VKT 3)",
    5: "Vavaka fangatahana",
}
DEPUIS_JAMAIS = 1000

def _dossier_assets():
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, "assets")


def charger_langues():
    with open(os.path.join(_dossier_assets(), "langues.json"), encoding="utf-8") as f:
        return json.load(f)


# Textes de la page mobile / application Android (assets/langues.json : français + malagasy « brouillon »).
# {g} = nom du groupe, {n} = nombre, {m} = nombre.
LANGUES = charger_langues()
TEXTES_PAGE = [(c["cle"], c["desc"], c["fr"]) for c in LANGUES["cles"]]      # (clé, description, texte français)
TEXTES_MG = {c["cle"]: c["mg"] for c in LANGUES["cles"]}                     # texte malagasy de départ


def date_texte(d, langue):
    mois = LANGUES["mois"][langue][d.month - 1]
    return f"{d.day}{'er' if (d.day == 1 and langue == 'fr') else ''} {mois} {d.year}"


def periode_texte(d1, d2, langue):
    modele = LANGUES.get("periode_modele", {}).get(langue) or LANGUES["periode_modele"]["fr"]
    return modele.replace("{a}", date_texte(d1, langue)).replace("{b}", date_texte(d2, langue))


def pack_page():
    """Paquet de langues embarqué dans la page : {textes:{fr:{},mg:{}}, jours:{}, mois:{}, ...}"""
    pack = {"textes": {"fr": {c["cle"]: c["fr"] for c in LANGUES["cles"]},
                       "mg": {c["cle"]: c["mg"] for c in LANGUES["cles"]}}}
    for k in ("jours", "mois", "jours_courts", "mois_courts"):
        pack[k] = LANGUES[k]
    return pack


# ------------------------------------------------------------------ dates
def s2d(s):
    return dt.date.fromisoformat(s)


def d2s(d):
    return d.isoformat()


def lire_date(texte):
    """Accepte JJ/MM/AAAA, J/M/AAAA ou AAAA-MM-JJ."""
    t = texte.strip()
    if "-" in t and len(t.split("-")[0]) == 4:
        return dt.date.fromisoformat(t)
    j, m, a = t.replace(".", "/").split("/")
    a = int(a)
    if a < 100:
        a += 2000
    return dt.date(a, int(m), int(j))


def prochain_dimanche(d):
    while d.weekday() != 6:
        d += dt.timedelta(days=1)
    return d


def fin_de_semestre(d):
    return dt.date(d.year, 6, 30) if d.month <= 6 else dt.date(d.year, 12, 31)


def est_dernier_dimanche_du_mois(d):
    return (d + dt.timedelta(days=7)).month != d.month


# ------------------------------------------------------------------ projet
def nouveau_projet(debut=None):
    debut = prochain_dimanche(debut or dt.date.today())
    return {
        "version": 1,
        "debut": d2s(debut),
        "fin": d2s(fin_de_semestre(debut)),
        "ordre": list(GROUPES_DEFAUT),
        "premier_groupe": GROUPES_DEFAUT[0],
        "tokana": True,
        "role5": False,
        "url_maj": "",
        "textes": {},
        "langue": "fr",
        "langue_libre": False,
        "groupes": {g: {"membres": []} for g in GROUPES_DEFAUT},
        "programme": [],
        "affectations": {},
        "verrous": [],
        "historique": {},
    }


def sauver(p, chemin):
    with open(chemin, "w", encoding="utf-8") as f:
        json.dump(p, f, ensure_ascii=False, indent=1)


def charger(chemin):
    with open(chemin, encoding="utf-8") as f:
        p = json.load(f)
    base = nouveau_projet()
    for k, v in base.items():
        p.setdefault(k, v)
    return p


def noms_du_groupe(p, g):
    return [m["nom"] for m in p["groupes"][g]["membres"]]


def ajouter_membre(p, g, nom):
    nom = nom.strip()
    if not nom:
        raise ValueError("Nom vide.")
    if nom in noms_du_groupe(p, g):
        raise ValueError(f"« {nom} » existe déjà dans le groupe {g}.")
    p["groupes"][g]["membres"].append({"nom": nom, "absences": []})


def renommer_membre(p, g, ancien, nouveau):
    nouveau = nouveau.strip()
    if not nouveau:
        raise ValueError("Nom vide.")
    if nouveau != ancien and nouveau in noms_du_groupe(p, g):
        raise ValueError(f"« {nouveau} » existe déjà dans le groupe {g}.")
    for m in p["groupes"][g]["membres"]:
        if m["nom"] == ancien:
            m["nom"] = nouveau
    for cle, aff in p["affectations"].items():
        if cle.endswith("|" + g):
            for r, n in aff.items():
                if n == ancien:
                    aff[r] = nouveau
    h = p["historique"].get(g, {})
    if ancien in h:
        h[nouveau] = h.pop(ancien)


def supprimer_membre(p, g, nom):
    p["groupes"][g]["membres"] = [m for m in p["groupes"][g]["membres"] if m["nom"] != nom]


def est_absent(membre, d):
    iso = d2s(d)
    return any(a <= iso <= b for a, b in membre.get("absences", []))


# ------------------------------------------------------------------ programme (qui passe quand)
def generer_programme(p):
    n = len(p["ordre"])
    if n == 0:
        raise ValueError("Aucun groupe.")
    if p["premier_groupe"] not in p["ordre"]:
        p["premier_groupe"] = p["ordre"][0]
    d = prochain_dimanche(s2d(p["debut"]))
    fin = s2d(p["fin"])
    idx = p["ordre"].index(p["premier_groupe"])
    prog = []
    while d <= fin:
        k = 1 if (p["tokana"] and est_dernier_dimanche_du_mois(d)) else 2
        k = min(k, n)
        prog.append({"date": d2s(d), "groupes": [p["ordre"][(idx + i) % n] for i in range(k)]})
        idx += k
        d += dt.timedelta(days=7)
    p["debut"] = d2s(prochain_dimanche(s2d(p["debut"])))
    p["programme"] = prog
    valides = {f"{it['date']}|{g}" for it in prog for g in it["groupes"]}
    p["affectations"] = {k: v for k, v in p["affectations"].items() if k in valides}
    p["verrous"] = [k for k in p["verrous"] if k in valides]


# ------------------------------------------------------------------ répartition équilibrée
def _etat_initial(p, g):
    hist = p.get("historique", {}).get(g, {})
    etat = {}
    for i, m in enumerate(p["groupes"][g]["membres"]):
        h = hist.get(m["nom"], {})
        roles = {r: int(h.get("roles", {}).get(str(r), 0)) for r in (1, 2, 3, 4, 5)}
        n = int(h.get("n", 0))
        etat[m["nom"]] = {
            "i": i, "n": n, "n0": n,
            "roles": dict(roles), "roles0": dict(roles),
            "dernier_role": int(h.get("dernier_role", 0)),
            "dernier": -(int(h.get("depuis", DEPUIS_JAMAIS)) + 1),
            "date_dernier": None,
            "absences": m.get("absences", []),
        }
    return etat


def _appliquer(e, role, j, date_iso):
    e["n"] += 1
    e["roles"][role] += 1
    e["dernier_role"] = role
    e["dernier"] = j
    e["date_dernier"] = date_iso


def _choisir_et_repartir(etat, d, j, role5):
    """Choisit les participants (les moins sollicités, les plus reposés) puis leur rôle
    (celui qu'ils ont le moins tenu, en évitant de répéter le dernier)."""
    besoin = 5 if role5 else 4
    dispo = [nm for nm, e in etat.items() if not est_absent(e, d)]
    dispo.sort(key=lambda nm: (etat[nm]["n"], -(j - etat[nm]["dernier"]), etat[nm]["i"]))
    choisis = dispo[:besoin]
    roles = [1, 2, 3, 4, 5][:min(len(choisis), besoin)]
    meilleur, cout_min = None, None
    for perm in itertools.permutations(roles, len(choisis)):
        cout = sum(etat[nm]["roles"][r] * 10 + (3 if etat[nm]["dernier_role"] == r else 0)
                   for nm, r in zip(choisis, perm))
        if cout_min is None or cout < cout_min:
            meilleur, cout_min = perm, cout
    return dict(zip(choisis, meilleur or ())), len(dispo)


def generer_affectations(p, garder_verrous=True):
    """Calcule qui prend quel rôle à chaque passage. Retourne la liste des avertissements."""
    avert = []
    if not garder_verrous:
        p["verrous"] = []
    verrous = set(p["verrous"])
    besoin = 5 if p["role5"] else 4
    nouvelles = {}
    etats = {g: _etat_initial(p, g) for g in p["groupes"]}
    compteur = {g: 0 for g in p["groupes"]}

    for it in p["programme"]:
        d = s2d(it["date"])
        for g in it["groupes"]:
            if g not in p["groupes"]:
                avert.append(f"{d:%d/%m/%Y} : groupe {g} inconnu.")
                continue
            etat, j = etats[g], compteur[g]
            compteur[g] += 1
            cle = f"{it['date']}|{g}"
            if cle in verrous and cle in p["affectations"]:
                aff = dict(p["affectations"][cle])
                for r, nom in aff.items():
                    if nom in etat:
                        _appliquer(etat[nom], int(r), j, it["date"])
                    elif nom:
                        avert.append(f"{d:%d/%m/%Y} {g} : « {nom} » n'est plus dans le groupe (choix manuel conservé).")
                nouvelles[cle] = aff
                continue
            attribution, nb_dispo = _choisir_et_repartir(etat, d, j, p["role5"])
            if nb_dispo < besoin:
                avert.append(f"{d:%d/%m/%Y} {g} : seulement {nb_dispo} membre(s) disponible(s) "
                             f"pour {besoin} rôles.")
            aff = {str(r): "" for r in range(1, besoin + 1)}
            for nom, r in attribution.items():
                aff[str(r)] = nom
                _appliquer(etat[nom], r, j, it["date"])
            nouvelles[cle] = aff
    p["affectations"] = nouvelles
    return avert


def verifier(p):
    """Contrôles sur les affectations actuelles (rôle vide, absent, doublon dans un passage)."""
    avert = []
    membres = {g: {m["nom"]: m for m in p["groupes"][g]["membres"]} for g in p["groupes"]}
    besoin = 5 if p["role5"] else 4
    for it in p["programme"]:
        d = s2d(it["date"])
        for g in it["groupes"]:
            aff = p["affectations"].get(f"{it['date']}|{g}", {})
            vus = []
            for r in range(1, besoin + 1):
                nom = aff.get(str(r), "")
                if not nom:
                    avert.append(f"{d:%d/%m/%Y} {g} : rôle {r} non attribué.")
                    continue
                if nom in vus:
                    avert.append(f"{d:%d/%m/%Y} {g} : « {nom} » a deux rôles.")
                vus.append(nom)
                m = membres[g].get(nom)
                if m is not None and est_absent(m, d):
                    avert.append(f"{d:%d/%m/%Y} {g} : « {nom} » est marqué(e) absent(e).")
    return avert


def _rejouer(p):
    """Rejoue les affectations actuelles (sans rien recalculer)."""
    etats = {g: _etat_initial(p, g) for g in p["groupes"]}
    compteur = {g: 0 for g in p["groupes"]}
    for it in p["programme"]:
        for g in it["groupes"]:
            if g not in etats:
                continue
            j = compteur[g]
            compteur[g] += 1
            aff = p["affectations"].get(f"{it['date']}|{g}", {})
            for r, nom in aff.items():
                if nom in etats[g]:
                    _appliquer(etats[g][nom], int(r), j, it["date"])
    return etats, compteur


def statistiques(p):
    """{groupe: [ {nom, periode, roles{1..5}, total, dernier} ]}"""
    etats, _ = _rejouer(p)
    out = {}
    for g, etat in etats.items():
        lignes = []
        for nom, e in etat.items():
            lignes.append({
                "nom": nom,
                "periode": e["n"] - e["n0"],
                "roles": {r: e["roles"][r] - e["roles0"][r] for r in (1, 2, 3, 4, 5)},
                "total": e["n"],
                "dernier": e["date_dernier"],
            })
        out[g] = lignes
    return out


# ------------------------------------------------------------------ période suivante
def continuer(p):
    """Nouveau projet qui démarre après le dernier dimanche de `p`, en gardant
    l'historique (qui a participé combien de fois, avec quel rôle) pour rester équitable."""
    if not p["programme"]:
        raise ValueError("Générez d'abord le programme.")
    etats, compteur = _rejouer(p)
    q = copy.deepcopy(p)
    hist = {}
    for g, etat in etats.items():
        jf = compteur[g]
        hist[g] = {}
        for nom, e in etat.items():
            hist[g][nom] = {
                "n": e["n"],
                "roles": {str(r): v for r, v in e["roles"].items()},
                "dernier_role": e["dernier_role"],
                "depuis": (jf - 1) - e["dernier"],
            }
    dernier = p["programme"][-1]
    debut = s2d(dernier["date"]) + dt.timedelta(days=7)
    idx = p["ordre"].index(dernier["groupes"][-1]) + 1
    q["historique"] = hist
    q["debut"] = d2s(debut)
    q["fin"] = d2s(fin_de_semestre(debut))
    q["premier_groupe"] = p["ordre"][idx % len(p["ordre"])]
    q["programme"], q["affectations"], q["verrous"] = [], {}, []
    for g in q["groupes"].values():
        for m in g["membres"]:
            m["absences"] = [a for a in m.get("absences", []) if a[1] >= q["debut"]]
    return q


# ------------------------------------------------------------------ sorties
def vers_dimanches(p):
    """Format attendu par le générateur Word : {date: [(groupe, [4 noms]), (groupe, [4 noms])]}"""
    out = {}
    for it in p["programme"]:
        bloc = []
        for g in it["groupes"]:
            aff = p["affectations"].get(f"{it['date']}|{g}", {})
            bloc.append((g, [aff.get(str(r), "") for r in (1, 2, 3, 4)]))
        while len(bloc) < 2:
            bloc.append(("", [""] * 4))
        out[s2d(it["date"])] = bloc
    return out


def exporter_excel(p, chemin):
    """Classeur au même format que 25-26.xlsx (feuille Fandaharana_2022) + feuille Equite."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Fandaharana_2022"
    libelles = ["Vavaka vakiteny", "Vakiteny I", "Vakiteny II", "Vakiteny III"]
    r = 1
    for it in p["programme"]:
        ws.cell(r, 1, s2d(it["date"])).number_format = "dd/mm/yyyy"
        ws.cell(r, 2, "GROUPE").font = Font(bold=True)
        for col, g in zip((4, 5), it["groupes"]):
            ws.cell(r, col, g).font = Font(bold=True)
            aff = p["affectations"].get(f"{it['date']}|{g}", {})
            for k in range(4):
                ws.cell(r + 1 + k, col, aff.get(str(k + 1), ""))
        for k, lib in enumerate(libelles):
            ws.cell(r + 1 + k, 2, lib)
        r += 6
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 20

    we = wb.create_sheet("Equite")
    we.append(["Groupe", "Nom", "Participations (période)", "VVK VKT (1)", "VKT 1 (2)", "VKT 2 (3)",
               "VKT 3 (4)", "Rôle 5", "Total avec historique", "Dernier passage"])
    for c in we[1]:
        c.font = Font(bold=True)
    for g, lignes in statistiques(p).items():
        for l in lignes:
            we.append([g, l["nom"], l["periode"], l["roles"][1], l["roles"][2], l["roles"][3],
                       l["roles"][4], l["roles"][5], l["total"],
                       s2d(l["dernier"]) if l["dernier"] else None])
    wb.save(chemin)


def importer_excel(chemin, avec_programme=False):
    """Crée un projet à partir d'un classeur existant (feuilles de groupe + Fandaharana_2022,
    ou simple classeur type 25-26.xlsx)."""
    import warnings
    warnings.filterwarnings("ignore")
    from openpyxl import load_workbook
    import excel_vers_word as ew

    dimanches, _ = ew.lire_donnees(chemin, True)
    if not dimanches:
        raise ValueError("Aucun dimanche trouvé dans ce fichier.")
    dates = sorted(dimanches)
    sequence = [g for d in dates for g, _ in dimanches[d] if g]
    ordre = []
    for g in sequence:
        if g not in ordre:
            ordre.append(g)

    wb = load_workbook(chemin, data_only=True)
    p = nouveau_projet(dates[0])
    p["ordre"] = ordre
    p["premier_groupe"] = ordre[0]
    p["tokana"] = any(sum(1 for g, _ in dimanches[d] if g) == 1 for d in dates)
    p["groupes"] = {}
    for g in ordre:
        noms = []
        if g in wb.sheetnames:
            ws = wb[g]
            for r in range(11, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if isinstance(v, str) and v.strip() and v.strip() not in noms:
                    noms.append(v.strip())
        else:
            for d in dates:
                for gg, ns in dimanches[d]:
                    if gg == g:
                        for n in ns:
                            if n and n not in noms:
                                noms.append(n)
        p["groupes"][g] = {"membres": [{"nom": n, "absences": []} for n in noms]}

    p["debut"] = d2s(dates[0])
    if avec_programme:
        p["fin"] = d2s(dates[-1])
        for d in dates:
            groupes = [g for g, _ in dimanches[d] if g]
            p["programme"].append({"date": d2s(d), "groupes": groupes})
            for g, ns in dimanches[d]:
                if not g:
                    continue
                cle = f"{d2s(d)}|{g}"
                p["affectations"][cle] = {str(i + 1): n for i, n in enumerate(ns)}
                p["verrous"].append(cle)
                for n in ns:
                    if n and n not in noms_du_groupe(p, g):
                        ajouter_membre(p, g, n)
    else:
        p["fin"] = d2s(fin_de_semestre(dates[0]))
    return p


# ------------------------------------------------------------------ page mobile (lecture seule)
def donnees_mobile(p):
    """Données publiées pour les membres (page mobile, application Android, mise à jour en ligne).
    "source" = adresse où retrouver la dernière version ; "maj" = date de cette version."""
    nb = 5 if p["role5"] else 4
    dimanches = []
    for it in p["programme"]:
        a = {}
        for g in it["groupes"]:
            aff = p["affectations"].get(f"{it['date']}|{g}", {})
            a[g] = [aff.get(str(r), "") for r in range(1, nb + 1)]
        dimanches.append({"d": it["date"], "g": it["groupes"], "a": a})
    if not dimanches:
        raise ValueError("Le programme est vide.")
    d1, d2 = s2d(dimanches[0]["d"]), s2d(dimanches[-1]["d"])
    periodes = {lg: periode_texte(d1, d2, lg) for lg in ("fr", "mg")}
    maintenant = dt.datetime.now(dt.timezone.utc)
    return {
        "periode": periodes["fr"],        # ancien format (pages plus anciennes)
        "periodes": periodes,             # période écrite dans chaque langue : la page choisit selon le membre
        "genere": dt.date.today().isoformat(),
        "maj": maintenant.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": (p.get("url_maj") or "").strip(),
        "langue": p.get("langue") if p.get("langue") in ("fr", "mg") else "fr",
        "langue_libre": bool(p.get("langue_libre", False)),     # le membre peut-il changer de langue ?
        "textes": {k: v for k, v in (p.get("textes") or {}).items() if isinstance(v, str) and v.strip()},
        "ordre": list(p["ordre"]),
        "roles": nb,
        "membres": {g: noms_du_groupe(p, g) for g in p["ordre"]},
        "dimanches": dimanches,
    }


def exporter_programme_json(p, chemin):
    """Fichier programme.json à déposer en ligne : les membres le reçoivent automatiquement."""
    d = donnees_mobile(p)
    with open(chemin, "w", encoding="utf-8") as f:
        f.write(json.dumps(d, ensure_ascii=False, separators=(",", ":")))


def construire_page(donnees, modele, logo_png, icone_png):
    """HTML complet de la page mobile : modèle + données + logo + paquet de langues."""
    import base64

    def datauri(f):
        with open(f, "rb") as fh:
            return "data:image/png;base64," + base64.b64encode(fh.read()).decode("ascii")

    def js(obj):
        t = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
        return t.replace("</", "<\\/").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")

    with open(modele, encoding="utf-8") as f:
        html = f.read()
    return (html.replace("__LANGUES__", js(pack_page())).replace("__DONNEES__", js(donnees))
            .replace("__LOGO__", datauri(logo_png)).replace("__ICONE__", datauri(icone_png)))


def exporter_mobile(p, chemin, modele, logo_png, icone_png):
    """Crée UNE page web autonome (HTML) à envoyer aux membres : chacun choisit son groupe et son nom,
    voit sa prochaine participation, son groupe, le programme d'un dimanche, et peut ajouter des rappels
    à son calendrier. Aucune modification possible côté membre."""
    html = construire_page(donnees_mobile(p), modele, logo_png, icone_png)
    with open(chemin, "w", encoding="utf-8") as f:
        f.write(html)
