# -*- coding: utf-8 -*-
"""
Page de téléchargement + QR code pour les membres.

Crée, dans un dossier, tout ce qu'il faut déposer sur un dépôt GitHub public (GitHub Pages) :
  - index.html            : la page (français / malagasy) avec un bouton par fichier ;
  - les fichiers à télécharger (copiés ici : l'application Android, la page du programme, des PDF...) ;
  - QR_telechargements.png: le QR code de l'adresse de la page (il ne change jamais) ;
  - Feuille_QR.docx       : une feuille à imprimer (QR + consignes en français et en malagasy).
"""
import datetime as dt
import json
import os
import re
import shutil
import urllib.parse

import fdl_core as core

TYPES = ("apk", "page", "autre")
LIBELLES_DEFAUT = {
    "apk": ("Application Android (rappels et programme)", "Rindranasa Android (fampahatsiahivana sy fandaharana)"),
    "page": ("Page du programme (tous les téléphones, navigateur)", "Pejy fandaharana (finday rehetra, navigateur)"),
}
TITRES = ("Téléchargements FDL", "Fisintonana FDL")
_MOTIF = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")


def url_page(utilisateur, depot):
    """Adresse de la page publiée avec GitHub Pages (dépôt public)."""
    return f"https://{utilisateur.strip().lower()}.github.io/{depot.strip()}/"


def nom_publie_propose(chemin, type_):
    """Nom stable du fichier en ligne (un APK s'appelle toujours FDL.apk, pour que le lien ne change jamais)."""
    base = os.path.basename(chemin)
    if type_ == "apk":
        return "FDL.apk"
    nettoye = re.sub(r"[^A-Za-z0-9._-]+", "_", base)
    return nettoye if _MOTIF.match(nettoye) else "fichier_" + nettoye.lstrip("._-")


def verifier(config):
    """Liste des problèmes à corriger avant de créer la page (vide = tout va bien)."""
    pb = []
    u, d = config.get("utilisateur", "").strip(), config.get("depot", "").strip()
    if not u or not _MOTIF.match(u):
        pb.append("Nom d'utilisateur GitHub manquant ou invalide.")
    if not d or not _MOTIF.match(d):
        pb.append("Nom du dépôt manquant ou invalide (lettres, chiffres, - _ . seulement).")
    fichiers = config.get("fichiers", [])
    if not fichiers:
        pb.append("Ajoutez au moins un fichier.")
    noms = []
    for f in fichiers:
        if not os.path.isfile(f.get("chemin", "")):
            pb.append(f"Fichier introuvable : {f.get('chemin', '')}")
        nom = f.get("nom") or nom_publie_propose(f.get("chemin", ""), f.get("type", "autre"))
        if nom in noms:
            pb.append(f"Deux fichiers portent le même nom en ligne : {nom}")
        noms.append(nom)
    return pb


def _date_texte(d, langue):
    return core.date_texte(d, langue)


def construire(config, dossier, modele, logo_png, icone_png):
    """Crée la page, copie les fichiers, fabrique le QR et la feuille à imprimer. Retourne un résumé."""
    import base64
    import segno

    pb = verifier(config)
    if pb:
        raise ValueError("\n".join(pb))
    os.makedirs(dossier, exist_ok=True)
    url = url_page(config["utilisateur"], config["depot"])

    fichiers = []
    for f in config["fichiers"]:
        type_ = f.get("type") if f.get("type") in TYPES else "autre"
        nom = f.get("nom") or nom_publie_propose(f["chemin"], type_)
        cible = os.path.join(dossier, nom)
        if os.path.abspath(f["chemin"]) != os.path.abspath(cible):
            shutil.copy2(f["chemin"], cible)
        defaut = LIBELLES_DEFAUT.get(type_, (os.path.splitext(nom)[0], os.path.splitext(nom)[0]))
        lib_fr = (f.get("libelle_fr") or "").strip() or defaut[0]
        lib_mg = (f.get("libelle_mg") or "").strip() or (f.get("libelle_fr") or "").strip() or defaut[1]
        fichiers.append({"nom": nom, "href": urllib.parse.quote(nom), "type": type_, "octets": os.path.getsize(cible),
                         "libelle_fr": lib_fr, "libelle_mg": lib_mg})

    aujourdhui = dt.date.today()
    donnees = {
        "langue_defaut": config.get("langue_defaut") if config.get("langue_defaut") in ("fr", "mg") else "fr",
        "langue_libre": bool(config.get("langue_libre", False)),      # le membre peut-il changer de langue ?
        "titre": {"fr": (config.get("titre_fr") or "").strip() or TITRES[0],
                  "mg": (config.get("titre_mg") or "").strip() or TITRES[1]},
        "maj": {"fr": _date_texte(aujourdhui, "fr"), "mg": _date_texte(aujourdhui, "mg")},
        "fichiers": fichiers,
    }

    def js(obj):
        t = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
        return t.replace("</", "<\\/").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")

    def datauri(f):
        with open(f, "rb") as fh:
            return "data:image/png;base64," + base64.b64encode(fh.read()).decode("ascii")

    with open(modele, encoding="utf-8") as fh:
        html = fh.read()
    html = html.replace("__DONNEES__", js(donnees)).replace("__LOGO__", datauri(logo_png)).replace("__ICONE__", datauri(icone_png))
    with open(os.path.join(dossier, "index.html"), "w", encoding="utf-8") as fh:
        fh.write(html)

    qr_png = os.path.join(dossier, "QR_telechargements.png")
    segno.make(url, error="m").save(qr_png, kind="png", scale=12, border=4)
    feuille = os.path.join(dossier, "Feuille_QR.docx")
    feuille_qr(feuille, url, qr_png, donnees["titre"]["fr"], donnees["titre"]["mg"])
    return {"url": url, "dossier": dossier, "qr": qr_png, "feuille": feuille, "fichiers": [x["nom"] for x in fichiers]}


def feuille_qr(chemin, url, qr_png, titre_fr, titre_mg):
    """Feuille A4 à imprimer : QR code + consignes en français et en malagasy."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    marine = RGBColor(0x17, 0x06, 0x7A)

    def ligne(texte, taille=12, gras=False, couleur=None, centre=True, apres=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(apres)
        r = p.add_run(texte)
        r.bold, r.font.size, r.font.name = gras, Pt(taille), "Calibri"
        if couleur is not None:
            r.font.color.rgb = couleur
        return p

    ligne("FJKM", 14, True, marine, apres=0)
    ligne(titre_fr, 26, True, marine, apres=0)
    ligne(titre_mg, 20, True, marine, apres=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(qr_png, width=Cm(10.5))
    ligne(url, 11, False, RGBColor(0x55, 0x55, 0x55), apres=14)
    ligne("Comment faire ?", 15, True, marine, centre=False, apres=3)
    for t in ("1. Ouvrez l'appareil photo de votre téléphone.", "2. Visez le QR code.",
              "3. Touchez le lien qui apparaît, puis choisissez ce que vous voulez télécharger."):
        ligne(t, 13, centre=False, apres=2)
    ligne("Ahoana ny fanaovana azy ?", 15, True, marine, centre=False, apres=3)
    for t in ("1. Sokafy ny appareil photo amin'ny findainao.", "2. Asehoy ny QR code.",
              "3. Tsindrio ny rohy miseho, dia fidio izay tianao hosintonina."):
        ligne(t, 13, centre=False, apres=2)
    doc.save(chemin)
