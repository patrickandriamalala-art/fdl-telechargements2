#!/usr/bin/env python3
"""
Génère le programme FDL (Word) à partir du planning Excel.

Usage :
    py excel_vers_word.py FDL_GP_5_Manu_Juillet-Dec_26.xlsm 2026 S2 programme.docx
        S1 = janvier-juin, S2 = juillet-décembre

Deux types de fichiers sont acceptés :
  A) le classeur source (.xlsm) avec les feuilles de groupe (ROTA, KEFASY, ISAIA,
     ESTERA, JAONA) : le script fait lui-même le travail du bouton UPDATE.
     Les dates de la ligne 10 des feuilles de groupe ne sont PAS nécessaires :
     le 1er passage d'un groupe dans Fandaharana_2022 correspond à la 1re colonne
     contenant des codes (1 à 4) dans sa feuille, le 2e passage à la colonne suivante, etc.
     Option --dates : retrouve plutôt la colonne grâce à la date écrite en ligne 10
     (comportement exact du bouton UPDATE d'origine).
  B) un classeur déjà rempli (type 25-26.xlsx, feuille Fandaharana_2022 complète).

Structure attendue de l'Excel (celle de 25-26.xlsx) :
  - colonne A : date du dimanche (vraie date), sur la ligne "GROUPE"
  - colonne B : GROUPE / Vavaka vakiteny / Vakiteny I / II / III
  - colonne D : Fidirana 1°   colonne E : Fidirana 2°
  - dernier dimanche du mois : un seul groupe (Fidirana tokana), en D ou E
"""
import sys
import datetime as dt
from openpyxl import load_workbook
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

MOIS = {1: "JANOARY", 2: "FEBROARY", 3: "MARTSA", 4: "APRILY", 5: "MEY", 6: "JONA",
        7: "JOLAY", 8: "AOGOSITRA", 9: "SEPTAMBRA", 10: "OKTOBRA", 11: "NOVAMBRA", 12: "DESAMBRA"}
LIGNES = ["VVK VKT", "VKT 1.", "VKT 2.", "VKT 3."]
GRIS = "D9D9D9"


# ---------------------------------------------------------------- lecture Excel
def clean(v):
    return str(v).strip() if v not in (None, "") else ""


def lire_excel(chemin):
    ws = load_workbook(chemin, data_only=True).worksheets[0]
    dimanches = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, dt.datetime) and clean(ws.cell(r, 2).value).upper() == "GROUPE":
            d = a.date()
            bloc = []
            for col in (4, 5):
                groupe = clean(ws.cell(r, col).value)
                noms = [clean(ws.cell(r + k, col).value) for k in range(1, 5)]
                bloc.append((groupe, noms))
            dimanches[d] = bloc
    return dimanches



CODES_ROLES = {1: "Vavaka vakiteny", 2: "Vakiteny I", 3: "Vakiteny II", 4: "Vakiteny III"}


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def lire_source(chemin, ordre=True):
    """Reproduit le bouton UPDATE : feuilles de groupe -> Fandaharana_2022.
    ordre=True  : le k-ième passage du groupe dans Fandaharana_2022 = la k-ième colonne
                  de codes de sa feuille, à partir de la 1re colonne qui contient des codes
                  (les dates de la ligne 10 ne servent pas).
    ordre=False : un groupe est rattaché à une date par la date de sa colonne (comme le VBA).
    Retourne (dimanches, avertissements) ou None si le fichier n'a pas ce format."""
    import warnings
    warnings.filterwarnings("ignore")
    wb = load_workbook(chemin, data_only=True)
    if "Fandaharana_2022" not in wb.sheetnames:
        return None
    fa = wb["Fandaharana_2022"]

    # 1) squelette : quel(s) groupe(s) à quelle date (colonnes D et E)
    squelette = []
    for r in range(1, fa.max_row + 1):
        d = _date(fa.cell(r, 1).value)
        if d and clean(fa.cell(r, 2).value).upper() == "GROUPE":
            squelette.append((d, clean(fa.cell(r, 4).value), clean(fa.cell(r, 5).value)))
    groupes_utiles = {g for _, g1, g2 in squelette for g in (g1, g2) if g}
    feuilles = [n for n in wb.sheetnames if n in groupes_utiles]
    if not feuilles:
        return None

    # 2) feuilles de groupe : dates en ligne 10, noms en colonne A, codes 0-5
    colonnes = {}      # groupe -> liste ordonnée de (date, {code: nom})
    avert = []
    for nom in feuilles:
        ws = wb[nom]
        def a_des_codes(c):
            return any(isinstance(ws.cell(r, c).value, (int, float)) and
                       int(ws.cell(r, c).value) in CODES_ROLES
                       for r in range(11, ws.max_row + 1))
        codees = [c for c in range(2, ws.max_column + 1) if a_des_codes(c)]
        cols = []
        if ordre:
            if codees:                      # de la 1re à la dernière colonne de codes
                cols = [(c, _date(ws.cell(10, c).value)) for c in range(codees[0], codees[-1] + 1)]
        else:
            for c in range(2, ws.max_column + 1):
                v = ws.cell(10, c).value
                d = _date(v)
                if d:
                    cols.append((c, d))
                elif v not in (None, "") and not isinstance(v, (int, float)):
                    avert.append(f"Feuille {nom} : date illisible en colonne {c} ({v!r}) "
                                 "- ouvrez et enregistrez le fichier dans Excel.")
        liste = []
        for c, d in cols:
            jour = {}
            for r in range(11, ws.max_row + 1):
                anarana = clean(ws.cell(r, 1).value)
                code = ws.cell(r, c).value
                if anarana and isinstance(code, (int, float)) and int(code) in CODES_ROLES:
                    jour[int(code)] = anarana          # le dernier gagne, comme le VBA
            liste.append((d, jour))
        colonnes[nom] = liste

    # 3) rattachement groupe/date -> colonne
    choix = {}   # (groupe, date) -> {code: nom}
    for g in feuilles:
        occ = [d for d, g1, g2 in squelette if g in (g1, g2)]
        dates_g = [d for d, _ in colonnes[g]]
        if not ordre:
            for d in occ:
                if d in dates_g:
                    choix[(g, d)] = colonnes[g][dates_g.index(d)][1]
        else:
            for j, d in enumerate(occ):
                if j < len(colonnes[g]):
                    choix[(g, d)] = colonnes[g][j][1]

    # 4) assemblage
    dimanches = {}
    for d, g1, g2 in squelette:
        bloc = []
        for g in (g1, g2):
            if not g:
                bloc.append(("", [""] * 4))
                continue
            jour = choix.get((g, d))
            if jour is None:
                avert.append(f"{d:%d/%m/%Y} : le groupe {g} est prévu mais la feuille {g} n'a pas de colonne pour ce passage (le groupe passe plus souvent que la feuille n'a de colonnes).")
                jour = {}
            noms = [jour.get(k, "") for k in (1, 2, 3, 4)]
            if not jour and (g, d) in choix:
                avert.append(f"{d:%d/%m/%Y} {g} : la colonne correspondante de la feuille {g} ne contient aucun code.")
            if jour:
                for k, n in zip((1, 2, 3, 4), noms):
                    if not n:
                        avert.append(f"{d:%d/%m/%Y} {g} : personne pour « {CODES_ROLES[k]} » (code {k}).")
            bloc.append((g, noms))
        dimanches[d] = bloc
    return dimanches, avert


def lire_donnees(chemin, ordre=True):
    r = lire_source(chemin, ordre)
    if r is not None:
        return r
    return lire_excel(chemin), []


# ---------------------------------------------------------------- helpers Word
def bordures(cell, on=True):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single" if on else "nil")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tcPr.append(b)


def gris(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), GRIS)
    tcPr.append(shd)


def vertical(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), "btLr")
    tcPr.append(td)


def nettoyer(cell):
    """Après fusion, supprime les paragraphes vides en trop (ils gonflent la hauteur)."""
    ps = cell.paragraphs
    for p in ps[1:]:
        if not p.text.strip():
            p._p.getparent().remove(p._p)


def ecrire(cell, texte, gras=False, taille=9, centre=False, vert_center=True):
    nettoyer(cell)
    p = cell.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texte)
    r.bold = gras
    r.font.size = Pt(taille)
    r.font.name = "Calibri"
    if vert_center:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def fixer_largeurs(table, largeurs):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), largeurs):
        gc.set(qn("w:w"), str(int(w * 567)))
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(largeurs) * 567)))
    for row in table.rows:
        for i, w in enumerate(largeurs):
            row.cells[i].width = Cm(w)


def marges_cellules(table):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, w in (("top", 0), ("left", 60), ("bottom", 0), ("right", 60)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


# ---------------------------------------------------------------- un tableau de 3 mois
def construire_tableau(doc, annee, mois3, dimanches):
    NB_BLOCS = 5
    W_LAB, W_GRP, W_NOM, W_SEP = 1.6, 0.65, 2.95, 0.45
    largeurs = []
    for i in range(3):
        largeurs += [W_LAB, W_GRP, W_NOM, W_GRP, W_NOM]
        if i < 2:
            largeurs.append(W_SEP)

    n_rows = 1 + NB_BLOCS * 5
    t = doc.add_table(rows=n_rows, cols=len(largeurs))
    marges_cellules(t)
    fixer_largeurs(t, largeurs)

    for r in range(n_rows):
        t.rows[r].height = Cm(0.62)
        trPr = t.rows[r]._tr.get_or_add_trPr()
        h = trPr.find(qn("w:trHeight"))
        h.set(qn("w:hRule"), "exact")

    for i, m in enumerate(mois3):
        c0 = i * 6                       # première colonne du mois
        # sundays du mois
        d = dt.date(annee, m, 1)
        while d.weekday() != 6:
            d += dt.timedelta(days=1)
        suns = []
        while d.month == m:
            suns.append(d)
            d += dt.timedelta(days=7)

        # en-tête du mois
        for c in range(c0, c0 + 5):
            bordures(t.cell(0, c))
        ent = t.cell(0, c0).merge(t.cell(0, c0 + 4))
        ecrire(ent, MOIS[m], gras=True, centre=True)

        for b in range(NB_BLOCS):
            r0 = 1 + b * 5
            if b >= len(suns):
                continue                 # pas de 5e dimanche : bloc vide, sans bordure
            s = suns[b]
            data = dimanches.get(s)
            n_actifs = sum(1 for g, _ in data if g) if data else 0
            # un seul groupe -> Fidirana tokana ; deux groupes -> blocs normaux même le dernier dimanche ;
            # aucun groupe connu -> le dernier dimanche du mois reste "tokana" comme dans le modèle
            tokana = (n_actifs == 1) or (n_actifs == 0 and b == len(suns) - 1)
            date_txt = f"{s.day}/{s.month}/{s.year % 100:02d}"

            # bordures sur toutes les cellules du bloc
            for r in range(r0, r0 + 5):
                for c in range(c0, c0 + 5):
                    bordures(t.cell(r, c))

            ecrire(t.cell(r0, c0), date_txt, gras=True)
            for k, lab in enumerate(LIGNES):
                ecrire(t.cell(r0 + 1 + k, c0), lab)

            if tokana:
                h = t.cell(r0, c0 + 1).merge(t.cell(r0, c0 + 4))
                ecrire(h, "FIDIRANA TOKANA", gras=True)
                groupe, noms = "", [""] * 4
                if data:
                    for g, n in data:
                        if g:
                            groupe, noms = g, n
                g_cell = t.cell(r0 + 1, c0 + 1).merge(t.cell(r0 + 4, c0 + 1))
                ecrire(g_cell, groupe, gras=True, taille=8, centre=True)
                vertical(g_cell)
                if not groupe:
                    gris(g_cell)
                for k in range(4):
                    n_cell = t.cell(r0 + 1 + k, c0 + 2).merge(t.cell(r0 + 1 + k, c0 + 4))
                    ecrire(n_cell, noms[k])
                    if not groupe:
                        gris(n_cell)
            else:
                for f, (col_g, col_n, titre) in enumerate(
                        ((c0 + 1, c0 + 2, "FIDIRANA 1°"), (c0 + 3, c0 + 4, "FIDIRANA 2°"))):
                    h = t.cell(r0, col_g).merge(t.cell(r0, col_n))
                    ecrire(h, titre, gras=True)
                    groupe, noms = ("", [""] * 4)
                    if data:
                        groupe, noms = data[f]
                    g_cell = t.cell(r0 + 1, col_g).merge(t.cell(r0 + 4, col_g))
                    ecrire(g_cell, groupe, gras=True, taille=8, centre=True)
                    vertical(g_cell)
                    if not groupe:
                        gris(g_cell)
                    for k in range(4):
                        c = t.cell(r0 + 1 + k, col_n)
                        ecrire(c, noms[k])
                        if not groupe:
                            gris(c)

    # colonnes séparatrices sans bordure
    for sep in (5, 11):
        for r in range(n_rows):
            bordures(t.cell(r, sep), on=False)
    return t


# ---------------------------------------------------------------- programme
def creer_word(dimanches, annee, semestre, sortie):
    """Écrit le programme Word. dimanches = {date: [(groupe1, [4 noms]), (groupe2, [4 noms])]}.
    Retourne la liste des dimanches du semestre absents de `dimanches`."""
    mois = list(range(1, 7)) if semestre.upper() == "S1" else list(range(7, 13))

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.0)
    sec.top_margin = sec.bottom_margin = Cm(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0

    for i, tri in enumerate((mois[:3], mois[3:])):
        if i:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"FDL {annee}  -  {MOIS[tri[0]].title()} / {MOIS[tri[-1]].title()}")
        r.bold = True
        r.font.size = Pt(12)
        construire_tableau(doc, annee, tri, dimanches)

    doc.save(sortie)

    manquants = []
    for m in mois:
        d = dt.date(annee, m, 1)
        while d.weekday() != 6:
            d += dt.timedelta(days=1)
        while d.month == m:
            if d not in dimanches:
                manquants.append(d.strftime("%d/%m/%Y"))
            d += dt.timedelta(days=7)
    return manquants


def main(xlsx, annee, semestre, sortie, ordre=True):
    dimanches, avert = lire_donnees(xlsx, ordre)
    mois = list(range(1, 7)) if semestre.upper() == "S1" else list(range(7, 13))
    manquants = creer_word(dimanches, annee, semestre, sortie)
    print("Fichier créé :", sortie)
    if avert:
        mois_set = set(mois)
        pertinents = [a for a in avert if not a[:2].isdigit() or int(a[3:5]) in mois_set and int(a[6:10]) == annee]
        if pertinents:
            print("À vérifier dans l'Excel :")
            for a in pertinents:
                print("  -", a)
    print("Dimanches sans données dans l'Excel :", manquants or "aucun")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a not in ("--ordre", "--dates")]
    if len(args) != 4:
        print(__doc__)
        sys.exit(1)
    main(args[0], int(args[1]), args[2], args[3], ordre="--dates" not in sys.argv)
